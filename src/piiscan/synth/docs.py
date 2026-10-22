import io
import struct
import zipfile
from email.message import EmailMessage

from .cfb import write_cfb


def text(lines: list[str], encoding: str = "utf-8") -> bytes:
    return "\n".join(lines).encode(encoding)


def pdf_text(lines: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    for line in lines:
        pdf.cell(0, 8, line, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def page_image(lines: list[str], rotate: float = 0):
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("L", (2550, 3300), 255)
    draw, font = ImageDraw.Draw(img), ImageFont.load_default(size=42)
    for i, line in enumerate(lines):
        draw.text((200, 250 + 100 * i), line, fill=0, font=font)
    return img.rotate(rotate, expand=True, fillcolor=255) if rotate else img


def filler_lines(rng, n: int = 24) -> list[str]:
    words = "service billing statement address payment internet plan date amount due customer notice".split()
    return [" ".join(rng.choice(words) for _ in range(rng.randint(4, 9))) for _ in range(n)]


def pdf_scanned(lines: list[str] = (), rotate: float = 0) -> bytes:
    buf = io.BytesIO()
    page_image(list(lines), rotate).save(buf, format="PDF", resolution=300)
    return buf.getvalue()


def pdf_form(fields: dict[str, str]) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, FloatObject, NameObject, TextStringObject

    w = PdfWriter()
    page = w.add_blank_page(612, 792)
    refs = ArrayObject()
    for i, (name, value) in enumerate(fields.items()):
        field = DictionaryObject({
            NameObject("/Type"): NameObject("/Annot"), NameObject("/Subtype"): NameObject("/Widget"),
            NameObject("/FT"): NameObject("/Tx"), NameObject("/T"): TextStringObject(name),
            NameObject("/V"): TextStringObject(value),
            NameObject("/Rect"): ArrayObject([FloatObject(x) for x in (72, 700 - 30 * i, 300, 720 - 30 * i)]),
        })
        refs.append(w._add_object(field))
    page[NameObject("/Annots")] = refs
    w._root_object[NameObject("/AcroForm")] = DictionaryObject({NameObject("/Fields"): refs})
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def pdf_encrypted(lines: list[str], password: str = "secret") -> bytes:
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_text(lines))))
    w.encrypt(user_password=password, algorithm="AES-256")
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def pdf_with_attachment(lines: list[str], name: str, data: bytes) -> bytes:
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter(clone_from=PdfReader(io.BytesIO(pdf_text(lines))))
    w.add_attachment(name, data)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def docx(paragraphs: list[str], header: str | None = None, comment: str | None = None, deleted: str | None = None,
         embedded: tuple[str, bytes] | None = None) -> bytes:
    from docx import Document

    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if header:
        d.sections[0].header.paragraphs[0].text = header
    if comment:
        run = d.add_paragraph("See note").runs[0]
        d.add_comment(run, text=comment, author="Reviewer")
    if deleted:
        d.add_paragraph("DELMARK")
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    if deleted:
        data = _rewrite_part(data, "word/document.xml",
                             lambda x: x.replace(b"<w:t>DELMARK</w:t>", f"<w:delText>{deleted}</w:delText>".encode()))
    if embedded:
        data = _add_part(data, f"word/embeddings/{embedded[0]}", embedded[1])
    return data


def xlsx(sheets: dict[str, list[list]], hidden: set[str] = frozenset()) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        if name in hidden:
            ws.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pptx(slide_text: str, notes: str | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = slide_text
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def xls(rows: list[list]) -> bytes:
    import xlwt

    wb = xlwt.Workbook()
    ws = wb.add_sheet("Accounts")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            ws.write(r, c, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def encrypted_ooxml(data: bytes, password: str = "secret") -> bytes:
    from msoffcrypto.format.ooxml import OOXMLFile

    out = io.BytesIO()
    OOXMLFile(io.BytesIO(data)).encrypt(password, out)
    return out.getvalue()


def doc_encrypted_stub() -> bytes:
    fib = bytearray(1472)
    struct.pack_into("<HH", fib, 0, 0xA5EC, 0x00C1)
    struct.pack_into("<H", fib, 0x0A, 0x0100 | 0x0200)
    return write_cfb({"WordDocument": bytes(fib), "1Table": b"\0" * 4096})


def rtf(body: str) -> bytes:
    escaped = body.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}").replace("\n", "\\par\n")
    return ("{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Times New Roman;}}\\f0 " + escaped + "}").encode("latin-1")


def eml(subject: str, body: str, attachments: list[tuple[str, bytes]] = (), html: str | None = None) -> bytes:
    m = EmailMessage()
    m["From"] = "billing@example.net"
    m["To"] = "customer@example.com"
    m["Subject"] = subject
    m.set_content(body)
    if html:
        m.add_alternative(html, subtype="html")
    for name, data in attachments:
        m.add_attachment(data, maintype="application", subtype="octet-stream", filename=name)
    return m.as_bytes()


def msg(subject: str, body: str, attachments: list[tuple[str, bytes]] = ()) -> bytes:
    def s(value):
        return value.encode("utf-16-le")

    strings = {0x001A: "IPM.Note", 0x0037: subject, 0x1000: body, 0x0C1A: "Billing Dept", 0x0E04: "Customer"}
    root = {f"__substg1.0_{pid:04X}001F": s(v) for pid, v in strings.items()}
    props = struct.pack("<8xIIII8x", 0, len(attachments), 0, len(attachments))
    props += b"".join(_var_prop(pid << 16 | 0x001F, len(s(v)) + 2) for pid, v in strings.items())
    root["__properties_version1.0"] = props
    root["__nameid_version1.0"] = {"__substg1.0_00020102": b"", "__substg1.0_00030102": b"", "__substg1.0_00040102": b""}
    for i, (name, data) in enumerate(attachments):
        root[f"__attach_version1.0_#{i:08X}"] = {
            "__substg1.0_3707001F": s(name), "__substg1.0_3704001F": s(name), "__substg1.0_37010102": data,
            "__properties_version1.0": b"\0" * 8 + _fixed_prop(0x37050003, 1) + _fixed_prop(0x0E210003, i)
            + _var_prop(0x3707001F, len(s(name)) + 2) + _var_prop(0x3704001F, len(s(name)) + 2)
            + _var_prop(0x37010102, len(data)),
        }
    return write_cfb(root)


def _fixed_prop(tag: int, value: int) -> bytes:
    return struct.pack("<IIi4x", tag, 0x6, value)


def _var_prop(tag: int, size: int) -> bytes:
    return struct.pack("<III4x", tag, 0x6, size)


def zip_(members: dict[str, bytes], encrypted: set[str] = frozenset()) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in members.items():
            z.writestr(name, data)
    data = bytearray(buf.getvalue())
    with zipfile.ZipFile(io.BytesIO(bytes(data))) as z:
        for info in z.infolist():
            if info.filename in encrypted:
                data[info.header_offset + 6] |= 0x1
                _set_cd_flag(data, info.filename)
    return bytes(data)


def _set_cd_flag(data: bytearray, filename: str) -> None:
    pos = 0
    while (pos := data.find(b"PK\x01\x02", pos)) != -1:
        name_len = struct.unpack_from("<H", data, pos + 28)[0]
        if data[pos + 46:pos + 46 + name_len] == filename.encode():
            data[pos + 8] |= 0x1
            return
        pos += 4


def tiff(pages: int = 0, images: list | None = None) -> bytes:
    from PIL import Image

    frames = images or [Image.new("L", (200, 100), 255) for _ in range(pages)]
    buf = io.BytesIO()
    frames[0].save(buf, format="TIFF", save_all=True, append_images=frames[1:])
    return buf.getvalue()


def heic() -> bytes:
    from PIL import Image
    from pillow_heif import register_heif_opener

    register_heif_opener()
    buf = io.BytesIO()
    Image.new("RGB", (64, 64), "white").save(buf, format="HEIF")
    return buf.getvalue()


def _rewrite_part(data: bytes, part: str, fn) -> bytes:
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            content = src.read(info)
            z.writestr(info, fn(content) if info.filename == part else content)
    return out.getvalue()


def _add_part(data: bytes, part: str, content: bytes) -> bytes:
    out = io.BytesIO(data)
    with zipfile.ZipFile(out, "a", zipfile.ZIP_DEFLATED) as z:
        z.writestr(part, content)
    return out.getvalue()
